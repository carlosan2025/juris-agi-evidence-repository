"""Background worker tasks for document processing."""

import hashlib
import logging
import mimetypes
import os
import tempfile
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
from rq import get_current_job
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

from evidence_repository.config import get_settings
from evidence_repository.models.document import Document, DocumentVersion, ExtractionStatus
from evidence_repository.queue.jobs import JobManager, JobType, get_job_manager
from evidence_repository.storage.local import LocalFilesystemStorage

logger = logging.getLogger(__name__)

# Supported MIME types mapping
MIME_TYPE_MAP = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


def _get_sync_db_session() -> Session:
    """Get synchronous database session for worker tasks.

    Workers run in a separate process and need sync connections.
    """
    settings = get_settings()
    # Convert async URL to sync
    sync_url = settings.database_url.replace("+asyncpg", "")
    engine = create_engine(sync_url)
    SessionLocal = sessionmaker(bind=engine)
    return SessionLocal()


def _get_storage() -> LocalFilesystemStorage:
    """Get storage backend for worker tasks."""
    settings = get_settings()
    return LocalFilesystemStorage(base_path=settings.file_storage_root)


def _update_progress(progress: float, message: str | None = None) -> None:
    """Update progress of the current job."""
    job = get_current_job()
    if job:
        job_manager = get_job_manager()
        job_manager.update_progress(job.id, progress, message)


def _compute_file_hash(data: bytes) -> str:
    """Compute SHA-256 hash of file content."""
    return hashlib.sha256(data).hexdigest()


# =============================================================================
# Document Processing Tasks
# =============================================================================


def task_ingest_document(
    file_data: bytes,
    filename: str,
    content_type: str,
    metadata: dict | None = None,
    user_id: str | None = None,
) -> dict:
    """Ingest a document into the repository.

    Args:
        file_data: File content as bytes.
        filename: Original filename.
        content_type: MIME type.
        metadata: Optional metadata.
        user_id: User who initiated the upload.

    Returns:
        Dict with document_id and version_id.
    """
    _update_progress(0, "Starting document ingestion")

    settings = get_settings()
    db = _get_sync_db_session()
    storage = _get_storage()

    try:
        file_hash = _compute_file_hash(file_data)
        _update_progress(10, "Computed file hash")

        # Check for duplicate
        existing = db.execute(
            select(Document).where(
                Document.file_hash == file_hash,
                Document.deleted_at.is_(None),
            )
        ).scalar_one_or_none()

        if existing:
            _update_progress(100, "Document already exists (deduplicated)")
            return {
                "document_id": str(existing.id),
                "version_id": str(existing.versions[0].id) if existing.versions else None,
                "deduplicated": True,
            }

        # Create document
        document = Document(
            filename=filename,
            original_filename=filename,
            content_type=content_type,
            file_hash=file_hash,
            metadata_=metadata or {},
        )
        db.add(document)
        db.flush()
        _update_progress(30, "Created document record")

        # Generate storage path
        version_number = 1
        storage_path = f"documents/{document.id}/v{version_number}/{filename}"

        # Upload to storage synchronously
        import asyncio
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            loop.run_until_complete(
                storage.upload(storage_path, file_data, content_type)
            )
        finally:
            loop.close()
        _update_progress(60, "Uploaded file to storage")

        # Create version
        version = DocumentVersion(
            document_id=document.id,
            version_number=version_number,
            storage_path=storage_path,
            file_size=len(file_data),
            file_hash=file_hash,
            extraction_status=ExtractionStatus.PENDING,
            metadata_={"uploaded_by": user_id} if user_id else {},
        )
        db.add(version)
        db.commit()
        _update_progress(100, "Document ingested successfully")

        return {
            "document_id": str(document.id),
            "version_id": str(version.id),
            "deduplicated": False,
        }

    except Exception as e:
        db.rollback()
        logger.error(f"Document ingestion failed: {e}")
        raise
    finally:
        db.close()


def task_extract_document(
    document_id: str,
    version_id: str | None = None,
) -> dict:
    """Extract text from a document.

    Args:
        document_id: Document ID.
        version_id: Optional specific version ID (defaults to latest).

    Returns:
        Dict with extraction results.
    """
    _update_progress(0, "Starting text extraction")

    db = _get_sync_db_session()
    storage = _get_storage()

    try:
        # Get document and version
        doc_uuid = uuid.UUID(document_id)
        document = db.execute(
            select(Document).where(Document.id == doc_uuid)
        ).scalar_one_or_none()

        if not document:
            raise ValueError(f"Document {document_id} not found")

        if version_id:
            version = db.execute(
                select(DocumentVersion).where(DocumentVersion.id == uuid.UUID(version_id))
            ).scalar_one_or_none()
        else:
            version = db.execute(
                select(DocumentVersion)
                .where(DocumentVersion.document_id == doc_uuid)
                .order_by(DocumentVersion.version_number.desc())
            ).scalars().first()

        if not version:
            raise ValueError(f"No version found for document {document_id}")

        _update_progress(10, "Found document version")

        # Update status
        version.extraction_status = ExtractionStatus.PROCESSING
        db.flush()

        # Download file content
        import asyncio
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            file_data = loop.run_until_complete(storage.download(version.storage_path))
        finally:
            loop.close()
        _update_progress(30, "Downloaded file from storage")

        # Extract based on content type
        text = ""
        page_count = None

        if document.content_type == "application/pdf":
            text, page_count = _extract_pdf_text(file_data)
        elif document.content_type in ["text/plain", "text/markdown", "text/csv"]:
            text = _extract_text_content(file_data)
        elif document.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            text = _extract_xlsx_text(file_data)
        elif document.content_type.startswith("image/"):
            text = _extract_image_text(file_data, document.content_type)
        else:
            raise ValueError(f"Unsupported content type: {document.content_type}")

        _update_progress(80, "Extracted text content")

        # Update version
        version.extracted_text = text
        version.extraction_status = ExtractionStatus.COMPLETED
        version.extracted_at = datetime.utcnow()
        version.page_count = page_count
        version.extraction_error = None
        db.commit()

        _update_progress(100, "Extraction completed successfully")

        return {
            "document_id": document_id,
            "version_id": str(version.id),
            "text_length": len(text),
            "page_count": page_count,
        }

    except Exception as e:
        if version:
            version.extraction_status = ExtractionStatus.FAILED
            version.extraction_error = str(e)
            db.commit()
        logger.error(f"Extraction failed for {document_id}: {e}")
        raise
    finally:
        db.close()


def task_embed_document(
    document_id: str,
    version_id: str | None = None,
) -> dict:
    """Generate embeddings for a document.

    Args:
        document_id: Document ID.
        version_id: Optional specific version ID.

    Returns:
        Dict with embedding results.
    """
    _update_progress(0, "Starting embedding generation")

    # Import here to avoid circular imports
    from evidence_repository.embeddings.chunker import TextChunker
    from evidence_repository.embeddings.openai_client import OpenAIEmbeddingClient
    from evidence_repository.models.embedding import EmbeddingChunk

    settings = get_settings()
    db = _get_sync_db_session()

    try:
        # Get version with extracted text
        doc_uuid = uuid.UUID(document_id)

        if version_id:
            version = db.execute(
                select(DocumentVersion).where(DocumentVersion.id == uuid.UUID(version_id))
            ).scalar_one_or_none()
        else:
            version = db.execute(
                select(DocumentVersion)
                .where(DocumentVersion.document_id == doc_uuid)
                .order_by(DocumentVersion.version_number.desc())
            ).scalars().first()

        if not version:
            raise ValueError(f"No version found for document {document_id}")

        if not version.extracted_text:
            raise ValueError(f"Document {document_id} has no extracted text")

        _update_progress(10, "Found document with extracted text")

        # Chunk the text
        chunker = TextChunker(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
        )
        chunks = chunker.chunk_text(version.extracted_text)
        _update_progress(20, f"Created {len(chunks)} text chunks")

        if not chunks:
            return {
                "document_id": document_id,
                "version_id": str(version.id),
                "chunks_created": 0,
            }

        # Generate embeddings using synchronous approach
        client = OpenAIEmbeddingClient()
        import asyncio
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)

        try:
            texts = [c.text for c in chunks]
            embeddings = loop.run_until_complete(client.embed_texts(texts))
        finally:
            loop.close()

        _update_progress(70, "Generated embeddings")

        # Delete existing embeddings for this version
        db.execute(
            EmbeddingChunk.__table__.delete().where(
                EmbeddingChunk.document_version_id == version.id
            )
        )

        # Store embedding chunks
        for chunk, embedding in zip(chunks, embeddings):
            embedding_chunk = EmbeddingChunk(
                document_version_id=version.id,
                chunk_index=chunk.index,
                text=chunk.text,
                embedding=embedding,
                char_start=chunk.char_start,
                char_end=chunk.char_end,
                metadata_=chunk.metadata,
            )
            db.add(embedding_chunk)

        db.commit()
        _update_progress(100, "Embeddings stored successfully")

        return {
            "document_id": document_id,
            "version_id": str(version.id),
            "chunks_created": len(chunks),
        }

    except Exception as e:
        db.rollback()
        logger.error(f"Embedding generation failed for {document_id}: {e}")
        raise
    finally:
        db.close()


def task_process_document_full(
    file_data: bytes,
    filename: str,
    content_type: str,
    metadata: dict | None = None,
    user_id: str | None = None,
    skip_embedding: bool = False,
) -> dict:
    """Full document processing pipeline: ingest -> extract -> embed.

    Args:
        file_data: File content.
        filename: Filename.
        content_type: MIME type.
        metadata: Optional metadata.
        user_id: User ID.
        skip_embedding: Skip embedding generation.

    Returns:
        Dict with all processing results.
    """
    result = {"steps": []}

    # Step 1: Ingest
    _update_progress(0, "Step 1/3: Ingesting document")
    ingest_result = task_ingest_document(
        file_data=file_data,
        filename=filename,
        content_type=content_type,
        metadata=metadata,
        user_id=user_id,
    )
    result["ingest"] = ingest_result
    result["steps"].append("ingest")

    # Step 2: Extract
    _update_progress(33, "Step 2/3: Extracting text")
    try:
        extract_result = task_extract_document(
            document_id=ingest_result["document_id"],
            version_id=ingest_result["version_id"],
        )
        result["extract"] = extract_result
        result["steps"].append("extract")
    except Exception as e:
        result["extract_error"] = str(e)
        logger.warning(f"Extraction failed, continuing: {e}")

    # Step 3: Embed (if extraction succeeded)
    if not skip_embedding and "extract" in result:
        _update_progress(66, "Step 3/3: Generating embeddings")
        try:
            embed_result = task_embed_document(
                document_id=ingest_result["document_id"],
                version_id=ingest_result["version_id"],
            )
            result["embed"] = embed_result
            result["steps"].append("embed")
        except Exception as e:
            result["embed_error"] = str(e)
            logger.warning(f"Embedding failed: {e}")

    _update_progress(100, "Processing complete")
    result["document_id"] = ingest_result["document_id"]
    result["version_id"] = ingest_result["version_id"]

    return result


# =============================================================================
# Bulk Ingestion Tasks
# =============================================================================


def task_bulk_folder_ingest(
    folder_path: str,
    recursive: bool = True,
    user_id: str | None = None,
    process_full: bool = True,
) -> dict:
    """Ingest all supported files from a folder.

    Args:
        folder_path: Path to folder to scan.
        recursive: Whether to scan subfolders.
        user_id: User initiating the bulk import.
        process_full: Whether to run full processing (extract + embed).

    Returns:
        Dict with results for each file.
    """
    settings = get_settings()
    supported_extensions = set(settings.supported_extensions)

    folder = Path(folder_path)
    if not folder.exists():
        raise ValueError(f"Folder not found: {folder_path}")
    if not folder.is_dir():
        raise ValueError(f"Not a directory: {folder_path}")

    _update_progress(0, f"Scanning folder: {folder_path}")

    # Find all supported files
    if recursive:
        files = [f for f in folder.rglob("*") if f.is_file() and f.suffix.lower() in supported_extensions]
    else:
        files = [f for f in folder.iterdir() if f.is_file() and f.suffix.lower() in supported_extensions]

    total_files = len(files)
    if total_files == 0:
        return {
            "folder_path": folder_path,
            "files_found": 0,
            "files_processed": 0,
            "results": [],
        }

    _update_progress(5, f"Found {total_files} files to process")

    results = []
    job_manager = get_job_manager()

    for i, file_path in enumerate(files):
        progress = 5 + (i / total_files) * 90
        _update_progress(progress, f"Processing {i + 1}/{total_files}: {file_path.name}")

        try:
            # Read file
            with open(file_path, "rb") as f:
                file_data = f.read()

            # Determine content type
            extension = file_path.suffix.lower()
            content_type = MIME_TYPE_MAP.get(extension, "application/octet-stream")

            # Process
            if process_full:
                result = task_process_document_full(
                    file_data=file_data,
                    filename=file_path.name,
                    content_type=content_type,
                    metadata={"source_path": str(file_path), "bulk_import": True},
                    user_id=user_id,
                    skip_embedding=False,
                )
            else:
                result = task_ingest_document(
                    file_data=file_data,
                    filename=file_path.name,
                    content_type=content_type,
                    metadata={"source_path": str(file_path), "bulk_import": True},
                    user_id=user_id,
                )

            results.append({
                "file_path": str(file_path),
                "status": "success",
                "result": result,
            })

        except Exception as e:
            logger.error(f"Failed to process {file_path}: {e}")
            results.append({
                "file_path": str(file_path),
                "status": "error",
                "error": str(e),
            })

    _update_progress(100, f"Completed processing {total_files} files")

    return {
        "folder_path": folder_path,
        "files_found": total_files,
        "files_processed": len([r for r in results if r["status"] == "success"]),
        "files_failed": len([r for r in results if r["status"] == "error"]),
        "results": results,
    }


def task_ingest_from_url(
    url: str,
    filename: str | None = None,
    user_id: str | None = None,
    process_full: bool = True,
) -> dict:
    """Download and ingest a file from a URL.

    Args:
        url: URL to download from.
        filename: Optional filename override.
        user_id: User ID.
        process_full: Whether to run full processing.

    Returns:
        Dict with ingestion results.
    """
    settings = get_settings()

    _update_progress(0, f"Downloading file from URL")

    try:
        # Download file
        with httpx.Client(timeout=settings.url_download_timeout) as client:
            response = client.get(url, follow_redirects=True)
            response.raise_for_status()

        file_data = response.content
        _update_progress(40, f"Downloaded {len(file_data)} bytes")

        # Determine filename
        if not filename:
            # Try to get from Content-Disposition header
            cd = response.headers.get("content-disposition", "")
            if "filename=" in cd:
                filename = cd.split("filename=")[-1].strip('"')
            else:
                # Extract from URL
                from urllib.parse import urlparse
                filename = Path(urlparse(url).path).name or "downloaded_file"

        # Determine content type
        content_type = response.headers.get("content-type", "").split(";")[0].strip()
        if not content_type or content_type == "application/octet-stream":
            # Guess from extension
            extension = Path(filename).suffix.lower()
            content_type = MIME_TYPE_MAP.get(extension, "application/octet-stream")

        # Check if supported
        extension = Path(filename).suffix.lower()
        if extension not in settings.supported_extensions:
            raise ValueError(f"Unsupported file type: {extension}")

        # Check file size
        max_size = settings.max_file_size_mb * 1024 * 1024
        if len(file_data) > max_size:
            raise ValueError(f"File too large: {len(file_data)} bytes (max {max_size})")

        _update_progress(50, "Processing downloaded file")

        # Process
        if process_full:
            result = task_process_document_full(
                file_data=file_data,
                filename=filename,
                content_type=content_type,
                metadata={"source_url": url, "url_import": True},
                user_id=user_id,
            )
        else:
            result = task_ingest_document(
                file_data=file_data,
                filename=filename,
                content_type=content_type,
                metadata={"source_url": url, "url_import": True},
                user_id=user_id,
            )

        _update_progress(100, "URL ingestion complete")

        return {
            "url": url,
            "filename": filename,
            "content_type": content_type,
            "file_size": len(file_data),
            "result": result,
        }

    except httpx.HTTPError as e:
        raise ValueError(f"Failed to download URL: {e}")


# =============================================================================
# Text Extraction Helpers
# =============================================================================


def _extract_pdf_text(data: bytes) -> tuple[str, int]:
    """Extract text from PDF using pypdf."""
    import pypdf
    import io

    reader = pypdf.PdfReader(io.BytesIO(data))
    text_parts = []

    for page in reader.pages:
        text = page.extract_text() or ""
        text_parts.append(text)

    return "\n\n".join(text_parts), len(reader.pages)


def _extract_text_content(data: bytes) -> str:
    """Extract text from plain text files."""
    for encoding in ["utf-8", "latin-1", "cp1252"]:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_xlsx_text(data: bytes) -> str:
    """Extract text from Excel files."""
    import openpyxl
    import io

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    text_parts = []

    for sheet in wb.worksheets:
        sheet_text = [f"=== Sheet: {sheet.title} ==="]
        for row in sheet.iter_rows(values_only=True):
            row_text = [str(cell) if cell is not None else "" for cell in row]
            if any(row_text):
                sheet_text.append("\t".join(row_text))
        text_parts.append("\n".join(sheet_text))

    return "\n\n".join(text_parts)


def _extract_image_text(data: bytes, content_type: str) -> str:
    """Extract text from images using OCR.

    Note: This is a placeholder. For real OCR, integrate Tesseract or a cloud OCR service.
    """
    # For now, just return a placeholder
    # In production, use pytesseract or cloud OCR
    return f"[Image content - {content_type}. OCR not implemented in this version.]"


# =============================================================================
# Batch Ingestion Tasks (with progress tracking)
# =============================================================================


def task_batch_folder_ingest(
    batch_id: str,
    folder_path: str,
    project_id: str | None = None,
    auto_process: bool = True,
    user_id: str | None = None,
) -> dict:
    """Process a folder ingestion batch with item-level tracking.

    This task processes each file in the batch, updating the ingestion_items
    table with progress and results.

    Args:
        batch_id: ID of the ingestion batch.
        folder_path: Path to the folder.
        project_id: Optional project to attach documents to.
        auto_process: Whether to auto-process documents.
        user_id: User initiating the ingestion.

    Returns:
        Dict with batch results.
    """
    from evidence_repository.models.ingestion import (
        IngestionBatch,
        IngestionBatchStatus,
        IngestionItem,
        IngestionItemStatus,
    )
    from evidence_repository.models.project import ProjectDocument

    db = _get_sync_db_session()
    storage = _get_storage()

    try:
        # Get batch
        batch_uuid = uuid.UUID(batch_id)
        batch = db.execute(
            select(IngestionBatch).where(IngestionBatch.id == batch_uuid)
        ).scalar_one_or_none()

        if not batch:
            raise ValueError(f"Batch {batch_id} not found")

        # Get all items
        items = db.execute(
            select(IngestionItem)
            .where(IngestionItem.batch_id == batch_uuid)
            .order_by(IngestionItem.created_at)
        ).scalars().all()

        total_items = len(items)
        if total_items == 0:
            batch.status = IngestionBatchStatus.COMPLETED
            batch.completed_at = datetime.utcnow()
            db.commit()
            return {"batch_id": batch_id, "items_processed": 0}

        # Update batch status
        batch.status = IngestionBatchStatus.PROCESSING
        batch.started_at = datetime.utcnow()
        db.flush()

        _update_progress(0, f"Processing {total_items} files")

        processed = 0
        successful = 0
        failed = 0
        skipped = 0

        for i, item in enumerate(items):
            progress = (i / total_items) * 100
            _update_progress(progress, f"Processing {i + 1}/{total_items}: {item.source_filename}")

            try:
                item.status = IngestionItemStatus.PROCESSING
                item.started_at = datetime.utcnow()
                item.attempts += 1
                db.flush()

                # Read file
                file_path = Path(item.source_path)
                if not file_path.exists():
                    item.status = IngestionItemStatus.FAILED
                    item.error_message = f"File not found: {item.source_path}"
                    item.error_code = "FILE_NOT_FOUND"
                    item.completed_at = datetime.utcnow()
                    failed += 1
                    db.flush()
                    continue

                with open(file_path, "rb") as f:
                    file_data = f.read()

                # Compute hash and check for duplicates
                file_hash = _compute_file_hash(file_data)
                item.file_hash = file_hash

                existing = db.execute(
                    select(Document).where(
                        Document.file_hash == file_hash,
                        Document.deleted_at.is_(None),
                    )
                ).scalar_one_or_none()

                if existing:
                    # Document exists, mark as skipped but still attach to project
                    item.status = IngestionItemStatus.SKIPPED
                    item.document_id = existing.id
                    item.document_version_id = existing.versions[0].id if existing.versions else None
                    item.completed_at = datetime.utcnow()
                    skipped += 1

                    # Attach to project if specified
                    if project_id:
                        _attach_to_project(db, existing.id, uuid.UUID(project_id), user_id)

                    db.flush()
                    continue

                # Determine content type
                extension = file_path.suffix.lower()
                content_type = item.content_type or MIME_TYPE_MAP.get(extension, "application/octet-stream")

                # Process the document
                if auto_process:
                    result = task_process_document_full(
                        file_data=file_data,
                        filename=item.source_filename,
                        content_type=content_type,
                        metadata={"source_path": item.source_path, "batch_id": batch_id},
                        user_id=user_id,
                    )
                else:
                    result = task_ingest_document(
                        file_data=file_data,
                        filename=item.source_filename,
                        content_type=content_type,
                        metadata={"source_path": item.source_path, "batch_id": batch_id},
                        user_id=user_id,
                    )

                # Update item with results
                item.status = IngestionItemStatus.COMPLETED
                item.document_id = uuid.UUID(result["document_id"])
                item.document_version_id = uuid.UUID(result["version_id"]) if result.get("version_id") else None
                item.completed_at = datetime.utcnow()
                successful += 1

                # Attach to project if specified
                if project_id:
                    _attach_to_project(db, item.document_id, uuid.UUID(project_id), user_id)

                db.flush()

            except Exception as e:
                logger.error(f"Failed to process item {item.id}: {e}")
                item.status = IngestionItemStatus.FAILED
                item.error_message = str(e)
                item.error_code = "PROCESSING_ERROR"
                item.completed_at = datetime.utcnow()
                failed += 1
                db.flush()

            processed += 1

            # Update batch progress
            batch.processed_items = processed
            batch.successful_items = successful
            batch.failed_items = failed
            batch.skipped_items = skipped
            db.flush()

        # Determine final batch status
        if failed == 0 and skipped == 0:
            batch.status = IngestionBatchStatus.COMPLETED
        elif successful == 0 and failed == total_items:
            batch.status = IngestionBatchStatus.FAILED
        else:
            batch.status = IngestionBatchStatus.PARTIAL

        batch.completed_at = datetime.utcnow()
        db.commit()

        _update_progress(100, f"Batch completed: {successful} succeeded, {failed} failed, {skipped} skipped")

        return {
            "batch_id": batch_id,
            "total_items": total_items,
            "successful": successful,
            "failed": failed,
            "skipped": skipped,
        }

    except Exception as e:
        db.rollback()
        logger.error(f"Batch folder ingestion failed: {e}")

        # Update batch status
        try:
            batch = db.execute(
                select(IngestionBatch).where(IngestionBatch.id == uuid.UUID(batch_id))
            ).scalar_one_or_none()
            if batch:
                batch.status = IngestionBatchStatus.FAILED
                batch.completed_at = datetime.utcnow()
                db.commit()
        except Exception:
            pass

        raise
    finally:
        db.close()


def task_batch_url_ingest(
    batch_id: str,
    item_id: str,
    url: str,
    filename: str | None = None,
    project_id: str | None = None,
    auto_process: bool = True,
    user_id: str | None = None,
) -> dict:
    """Process a URL ingestion with item-level tracking.

    Args:
        batch_id: ID of the ingestion batch.
        item_id: ID of the ingestion item.
        url: URL to download from.
        filename: Optional filename override.
        project_id: Optional project to attach to.
        auto_process: Whether to auto-process.
        user_id: User ID.

    Returns:
        Dict with ingestion results.
    """
    from evidence_repository.models.ingestion import (
        IngestionBatch,
        IngestionBatchStatus,
        IngestionItem,
        IngestionItemStatus,
    )
    from evidence_repository.utils.security import validate_url_for_ssrf, SSRFProtectionError

    settings = get_settings()
    db = _get_sync_db_session()

    try:
        # Get batch and item
        batch_uuid = uuid.UUID(batch_id)
        item_uuid = uuid.UUID(item_id)

        batch = db.execute(
            select(IngestionBatch).where(IngestionBatch.id == batch_uuid)
        ).scalar_one_or_none()

        item = db.execute(
            select(IngestionItem).where(IngestionItem.id == item_uuid)
        ).scalar_one_or_none()

        if not batch or not item:
            raise ValueError(f"Batch {batch_id} or item {item_id} not found")

        # Update statuses
        batch.status = IngestionBatchStatus.PROCESSING
        batch.started_at = datetime.utcnow()
        item.status = IngestionItemStatus.DOWNLOADING
        item.started_at = datetime.utcnow()
        item.attempts += 1
        db.flush()

        _update_progress(0, "Validating URL")

        # SSRF protection
        try:
            validate_url_for_ssrf(url)
        except SSRFProtectionError as e:
            item.status = IngestionItemStatus.FAILED
            item.error_message = f"URL validation failed: {e}"
            item.error_code = "SSRF_BLOCKED"
            item.completed_at = datetime.utcnow()
            batch.status = IngestionBatchStatus.FAILED
            batch.failed_items = 1
            batch.processed_items = 1
            batch.completed_at = datetime.utcnow()
            db.commit()
            raise ValueError(f"URL validation failed: {e}")

        _update_progress(10, "Downloading file")

        # Download file
        try:
            with httpx.Client(timeout=settings.url_download_timeout) as client:
                response = client.get(url, follow_redirects=True)
                response.raise_for_status()

            file_data = response.content
        except httpx.HTTPError as e:
            item.status = IngestionItemStatus.FAILED
            item.error_message = f"Download failed: {e}"
            item.error_code = "DOWNLOAD_ERROR"
            item.completed_at = datetime.utcnow()
            batch.status = IngestionBatchStatus.FAILED
            batch.failed_items = 1
            batch.processed_items = 1
            batch.completed_at = datetime.utcnow()
            db.commit()
            raise ValueError(f"Download failed: {e}")

        item.source_size = len(file_data)
        item.status = IngestionItemStatus.PROCESSING
        db.flush()

        _update_progress(40, f"Downloaded {len(file_data)} bytes")

        # Check file size
        max_size = settings.max_file_size_mb * 1024 * 1024
        if len(file_data) > max_size:
            item.status = IngestionItemStatus.FAILED
            item.error_message = f"File too large: {len(file_data)} bytes (max {max_size})"
            item.error_code = "FILE_TOO_LARGE"
            item.completed_at = datetime.utcnow()
            batch.status = IngestionBatchStatus.FAILED
            batch.failed_items = 1
            batch.processed_items = 1
            batch.completed_at = datetime.utcnow()
            db.commit()
            raise ValueError(f"File too large: {len(file_data)} bytes")

        # Determine filename if not provided
        if not filename:
            cd = response.headers.get("content-disposition", "")
            if "filename=" in cd:
                filename = cd.split("filename=")[-1].strip('"')
            else:
                from urllib.parse import urlparse
                filename = Path(urlparse(url).path).name or "downloaded_file"

        item.source_filename = filename

        # Determine content type
        content_type = response.headers.get("content-type", "").split(";")[0].strip()
        if not content_type or content_type == "application/octet-stream":
            extension = Path(filename).suffix.lower()
            content_type = MIME_TYPE_MAP.get(extension, "application/octet-stream")

        item.content_type = content_type

        # Check if supported
        extension = Path(filename).suffix.lower()
        if extension not in settings.supported_extensions:
            item.status = IngestionItemStatus.FAILED
            item.error_message = f"Unsupported file type: {extension}"
            item.error_code = "UNSUPPORTED_TYPE"
            item.completed_at = datetime.utcnow()
            batch.status = IngestionBatchStatus.FAILED
            batch.failed_items = 1
            batch.processed_items = 1
            batch.completed_at = datetime.utcnow()
            db.commit()
            raise ValueError(f"Unsupported file type: {extension}")

        # Compute hash
        file_hash = _compute_file_hash(file_data)
        item.file_hash = file_hash
        db.flush()

        _update_progress(50, "Processing document")

        # Process
        if auto_process:
            result = task_process_document_full(
                file_data=file_data,
                filename=filename,
                content_type=content_type,
                metadata={"source_url": url, "batch_id": batch_id},
                user_id=user_id,
            )
        else:
            result = task_ingest_document(
                file_data=file_data,
                filename=filename,
                content_type=content_type,
                metadata={"source_url": url, "batch_id": batch_id},
                user_id=user_id,
            )

        # Update item with results
        item.status = IngestionItemStatus.COMPLETED
        item.document_id = uuid.UUID(result["document_id"])
        item.document_version_id = uuid.UUID(result["version_id"]) if result.get("version_id") else None
        item.completed_at = datetime.utcnow()

        # Attach to project if specified
        if project_id:
            _attach_to_project(db, item.document_id, uuid.UUID(project_id), user_id)

        # Update batch
        batch.status = IngestionBatchStatus.COMPLETED
        batch.successful_items = 1
        batch.processed_items = 1
        batch.completed_at = datetime.utcnow()
        db.commit()

        _update_progress(100, "URL ingestion complete")

        return {
            "batch_id": batch_id,
            "item_id": item_id,
            "url": url,
            "filename": filename,
            "document_id": result["document_id"],
            "version_id": result.get("version_id"),
        }

    except Exception as e:
        db.rollback()
        logger.error(f"URL ingestion failed: {e}")

        # Update statuses
        try:
            batch = db.execute(
                select(IngestionBatch).where(IngestionBatch.id == uuid.UUID(batch_id))
            ).scalar_one_or_none()
            item = db.execute(
                select(IngestionItem).where(IngestionItem.id == uuid.UUID(item_id))
            ).scalar_one_or_none()

            if item and item.status != IngestionItemStatus.FAILED:
                item.status = IngestionItemStatus.FAILED
                item.error_message = str(e)
                item.completed_at = datetime.utcnow()

            if batch and batch.status != IngestionBatchStatus.FAILED:
                batch.status = IngestionBatchStatus.FAILED
                batch.failed_items = 1
                batch.processed_items = 1
                batch.completed_at = datetime.utcnow()

            db.commit()
        except Exception:
            pass

        raise
    finally:
        db.close()


def _attach_to_project(
    db: Session,
    document_id: uuid.UUID,
    project_id: uuid.UUID,
    user_id: str | None,
) -> None:
    """Attach a document to a project if not already attached."""
    from evidence_repository.models.project import ProjectDocument

    # Check if already attached
    existing = db.execute(
        select(ProjectDocument).where(
            ProjectDocument.project_id == project_id,
            ProjectDocument.document_id == document_id,
        )
    ).scalar_one_or_none()

    if not existing:
        project_doc = ProjectDocument(
            project_id=project_id,
            document_id=document_id,
            attached_by=user_id,
        )
        db.add(project_doc)
        db.flush()
